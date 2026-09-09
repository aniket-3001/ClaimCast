# -*- coding: utf-8 -*-
"""Direction C - Poster. Rockwell + Segoe UI, full-bleed blue panels, huge numerals."""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import *
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = os.path.join(ROOT, 'Team Rocket_IIIT Delhi - C Poster.pptx')
set_fonts("Rockwell", "Segoe UI")
R, S = "Rockwell", "Segoe UI"
PALEB = RGBColor(0xA9, 0xCB, 0xEA)
# grey reads as unfinished; these are the same neutrals with blue in them
TITLE = RGBColor(0x14, 0x18, 0x1C)      # slide headings
BODY = RGBColor(0x23, 0x28, 0x2D)       # all running text, near-black
MUTED = BODY                            # no faded tier; size and weight rank instead
SOFTB = RGBColor(0xC4, 0xDC, 0xEF)
ROSE = RGBColor(0xFF, 0x6B, 0x7F)
prs, sl = open_deck()
s1, s2, s3, s4, s5 = sl
W = 13.333

# the template's own footers: date is TextShape 4 on each slide,
# "-Public-" is TextBox 7 sitting on the slide master
for _sl in sl:
    drop(_sl, "TextShape 4")
    _pn = shape_named(_sl, "TextShape 3")          # page number, template grey
    if _pn is not None:
        # it is an <a:fld>, so it has no runs; paint every rPr in the body
        for _rpr in _pn.text_frame._txBody.iter():
            if _rpr.tag.endswith("}rPr") or _rpr.tag.endswith("}defRPr"):
                for _old in _rpr.findall(qn("a:solidFill")):
                    _rpr.remove(_old)
                _fill = _rpr.makeelement(qn("a:solidFill"), {})
                _clr = _rpr.makeelement(qn("a:srgbClr"), {"val": "005EB8"})
                _fill.append(_clr)
                _rpr.insert(0, _fill)
for _sh in list(prs.slide_master.shapes):
    if _sh.name in ("TextBox 7", "Date Placeholder 3", "Footer Placeholder 4"):
        _sh._element.getparent().remove(_sh._element)
for _lay in prs.slide_layouts:
    for _sh in list(_lay.shapes):
        if _sh.name.startswith(("Date Placeholder", "Footer Placeholder")):
            _sh._element.getparent().remove(_sh._element)

# ══════════════════════════════════════════════════════════════ 1 · TITLE
band = rect(s1, 0, 0, W, 4.02, BLUE)
send_to_back(band)
t = s1.shapes.title
t.left, t.top, t.width, t.height = Inches(0.9), Inches(0.62), Inches(11.53), Inches(0.34)
tf = t.text_frame
tf.clear()
para(tf, True, [("PRECISION CARE CHALLENGE 2026   ·   PHASE 1, IDEA SUBMISSION", 11, True, PALEB, False, S)],
     PP_ALIGN.CENTER)
sub = [p for p in s1.placeholders if p.placeholder_format.idx == 1][0]
sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(1.36), Inches(11.53), Inches(1.30)
tf = sub.text_frame
tf.clear()
para(tf, True, [("ClaimCast", 72, True, WHITE, False, R)], PP_ALIGN.CENTER)
rect(s1, 6.17, 2.86, 0.98, 0.045, ROSE)
text(s1, 1.4, 3.10, 10.53, 0.72,
     [{"runs": [("Know what your hospital bill will actually cost you. ", 18, False, SOFTB, False, S),
                ("Before", 18, True, WHITE, False, S), (" you are admitted, not ", 18, False, SOFTB, False, S),
                ("after", 18, True, WHITE, False, S), (" you are discharged.", 18, False, SOFTB, False, S)],
       "align": PP_ALIGN.CENTER, "line": 1.25}])
text(s1, 0.9, 4.44, 11.53, 0.34,
     [{"runs": [("TRACK: HOSPITALITY", 11, True, RED, False, S),
                ("   ·   Holistic Optimization System for "
                 "Policy-Integrated Admission & Treatment Intelligence",
                 12, False, BODY, False, S)], "align": PP_ALIGN.CENTER}])
rect(s1, 5.42, 5.20, 2.50, 0.035, INK)
text(s1, 0.9, 5.56, 11.53, 0.95,
     [{"runs": [("Team Rocket", 26, True, INK, False, R)], "align": PP_ALIGN.CENTER, "sa": 5},
      {"runs": [("Indraprastha Institute of Information Technology, Delhi (IIIT‑Delhi)",
                 12.5, False, BODY, False, S)], "align": PP_ALIGN.CENTER}])

# ══════════════════════════════════════════════════════════════ 2 · PROBLEM
retitle(s2, "Explain The Problem You Are Attempting to Solve", font=R, color=TITLE)
content_frame(s2, keep=False)
rect(s2, 0, 1.08, 4.71, 6.42, BLUE)
text(s2, 0.54, 1.46, 3.72, 1.66,
     [{"runs": [("India’s most expensive healthcare decision is made in ninety seconds, at an "
                 "admission counter, by someone who has never read their policy.", 17, True, WHITE, False, R)],
       "line": 1.26}])
rect(s2, 0.54, 3.36, 0.56, 0.032, RGBColor(0x6F, 0xA8, 0xDA))
text(s2, 0.54, 3.58, 3.72, 0.95,
     [{"runs": [("IRDAI’s chairman has said claim payouts often fall short of "
                 "what policyholders expect.",
                 14.5, False, RGBColor(0xE4, 0xEF, 0xF8), True, S)], "line": 1.35}])
text(s2, 0.54, 4.62, 3.72, 0.26, [{"runs": [("AJAY SETH, CHAIRMAN, IRDAI  ·  NOV 2025", 10.5, True, WHITE, False, R)]}])
text(s2, 0.54, 4.92, 3.72, 0.56,
     [{"runs": [("The regulator has named the gap. Policy bots explain the wording; none of "
                 "them price the decision in front of you.", 11, False, SOFTB, False, S)], "line": 1.35}])
text(s2, 0.54, 6.92, 3.9, 0.24,
     [{"runs": [("Proportionate deduction  ·  IRDAI/HLT/REG/CIR/151/06/2020  ·  maximal "
                 "case; scope varies by policy wording", 8.5, False,
                 RGBColor(0xA9, 0xCB, 0xEA), False, S)]}])

PX = 5.21
text(s2, PX, 1.38, 7.6, 0.26, [{"runs": [("THE ROOM-RENT TRAP", 10.5, True, RED, False, R)]}])
text(s2, PX, 1.62, 3.6, 1.36, [{"runs": [("50%", 108, True, RED, False, R)]}])
text(s2, PX + 3.06, 1.96, 4.0, 0.80,
     [{"runs": [("of the bill", 22, True, INK, False, R)], "sa": 0, "line": 1.14},
      {"runs": [("can be refused", 22, True, INK, False, R)], "line": 1.14}])
text(s2, PX, 3.16, 7.5, 0.80,
     [{"runs": [("Your policy caps the room at ", 13.5, False, INK, False, S), ("₹5,000", 13.5, True, INK, False, S),
                (" a day. The hospital gave you one at ", 13.5, False, INK, False, S),
                ("₹10,000", 13.5, True, INK, False, S),
                (". Surgeon, OT, anaesthesia, nursing: ", 13.5, False, INK, False, S),
                ("every charge your policy ties to room category", 13.5, True, INK, False, S),
                (" is scaled down by that same ratio.", 13.5, False, INK, False, S)], "line": 1.42}])
hline(s2, PX, 4.16, 7.62, HAIR)
for i, (big, cap, sub_) in enumerate([("₹39,085", "mean out-of-pocket per cancer hospitalisation",
                                        "12,148-patient multi-centre Indian study"),
                                       ("₹3.31 L", "out-of-pocket per cancer patient, per year",
                                        "diagnostics 36%  ·  medicines 45%")]):
    xx = PX + i * 3.90
    text(s2, xx, 4.36, 3.6, 0.52, [{"runs": [(big, 32, True, BLUE, False, R)]}])
    text(s2, xx, 4.94, 3.55, 0.40, [{"runs": [(cap, 11.5, False, INK, False, S)], "line": 1.3}])
    text(s2, xx, 5.32, 3.55, 0.24, [{"runs": [(sub_, 9.5, False, BODY, False, S)]}])
rect(s2, PX, 5.92, 0.07, 0.62, RED)
text(s2, PX + 0.28, 5.96, 7.30, 0.56,
     [{"runs": [("Cost shock drives distress financing, delayed care and outright ", 13, False, INK, False, S),
                ("treatment abandonment", 13, True, RED, False, S),
                (". It decides whether treatment finishes.", 13, False, INK, False, S)], "line": 1.3}])

# ══════════════════════════════════════════════════════════════ 3 · IDEA
retitle(s3, "Briefly Explain Your Idea/Solution", font=R, color=TITLE)
content_frame(s3, keep=False)
rect(s3, 0, 1.08, W, 1.44, BLUE)
text(s3, 0.54, 1.22, 12.25, 0.74,
     [{"runs": [("Insurers compute what you are owed ", 21, True, WHITE, False, R),
                ("after discharge", 21, True, ROSE, False, R),
                (".  ClaimCast computes it ", 21, True, WHITE, False, R),
                ("before admission", 21, True, RGBColor(0x9B, 0xE3, 0xC4), False, R), (".", 21, True, WHITE, False, R)]}])
text(s3, 0.54, 2.04, 12.25, 0.28,
     [{"runs": [("Same arithmetic. Different timestamp. That is the whole idea.", 13, False, SOFTB, False, S)]}])

forks = [("ADMISSION", "Which hospital", "network or not", "₹1.2 L"),
         ("ADMISSION", "Which room class", "proportionate deduction", "₹78,400"),
         ("INVESTIGATION", "Cashless or reimburse", "what you must float", "₹2 L"),
         ("PROCEDURE", "Implant / consumable", "sub-limits & List I", "₹45,000"),
         ("RECOVERY", "24-hour admission rule", "day-care vs in-patient", "₹22,000")]
fw, gap = 2.35, 0.15
FX0 = 0.54
for i, (stage, t1, t2, money) in enumerate(forks):
    x = FX0 + i * (fw + gap)
    rect(s3, x, 2.80, fw, 0.055, BLUE)
    text(s3, x, 2.94, fw, 0.26,
         [{"runs": [("%d · %s" % (i + 1, stage), 9, True, BLUE, False, S)]}])
    text(s3, x, 3.26, fw, 0.62,
         [{"runs": [(t1, 13, True, INK, False, R)], "sa": 2, "line": 1.1},
          {"runs": [(t2, 9.5, False, BODY, False, S)]}])
    text(s3, x, 3.94, fw, 0.32,
         [{"runs": [(money, 18, True, RED, False, R)], "sa": 0},
          {"runs": [("at stake", 9, False, BODY, False, S)]}])
text(s3, 0.54, 4.78, 12.25, 0.24,
     [{"runs": [("THE CARE JOURNEY THE PROBLEM STATEMENT NAMES", 9, True, BLUE, False, R),
                ("      we price the decision at every stage, and each fork closes permanently "
                 "the moment it passes", 9.5, False, BODY, False, S)], "align": PP_ALIGN.CENTER}])

rect(s3, 0.54, 5.16, 6.06, 1.28, RGBColor(0xF0, 0xF7, 0xF3))
rect(s3, 0.54, 5.16, 0.07, 1.28, GREEN)
text(s3, 0.86, 5.34, 5.5, 0.24, [{"runs": [("MATERIALITY FILTER  ·  OUR ANSWER TO FEATURE BLOAT", 9.5, True, GREEN, False, R)]}])
text(s3, 0.86, 5.62, 5.5, 0.72,
     [{"runs": [("We compute the money at stake and ", 12, False, INK, False, S),
                ("stay silent unless it clears ₹10,000", 12, True, GREEN, False, S),
                (". A caregiver at 2 a.m. sees only what is worth their attention.", 12, False, INK, False, S)],
       "line": 1.3}])
rect(s3, 6.86, 5.16, 5.93, 1.28, RGBColor(0xEC, 0xF2, 0xF8))
rect(s3, 6.86, 5.16, 0.07, 1.28, BLUE)
text(s3, 7.18, 5.34, 5.4, 0.24, [{"runs": [("IS THIS UNIQUE, OR AN IMPROVISATION?", 9.5, True, BLUE, False, R)]}])
text(s3, 7.18, 5.62, 5.4, 0.72,
     [{"runs": [("Unique. ", 12, True, BLUE, False, S), ("Existing tools ", 12, False, INK, False, S),
                ("retrieve", 12, True, INK, False, S), (" policy text. ClaimCast ", 12, False, INK, False, S),
                ("simulates the claim and optimises the decision", 12, True, INK, False, S),
                (".  Retrieval → Simulation.", 12, False, INK, False, S)], "line": 1.3}])
text(s3, 0.54, 6.62, 12.25, 0.30,
     [{"runs": [("YOU WALK IN KNOWING WHAT YOU WILL OWE.    ", 10.5, True, BLUE, False, R),
                ("Same admission, ₹78,400 less out of pocket, because the caregiver was told "
                 "before the decision closed.", 10.5, False, BODY, False, S)]}])

# ══════════════════════════════════════════════════════════════ 4 · SOLUTION
retitle(s4, "Explain Your Solution", font=R, color=TITLE)
content_frame(s4, keep=False)
rect(s4, 0, 1.08, W, 1.32, BLUE)
text(s4, 0.54, 1.26, 9.2, 0.72,
     [{"runs": [("We cannot tell you exactly what your bill will be. We can tell you exactly "
                 "how much of it your policy will refuse.", 19, True, WHITE, False, R)], "line": 1.22}])
text(s4, 10.02, 1.42, 2.78, 0.80,
     [{"runs": [("AI where uncertain.", 12.5, True, RGBColor(0x9B, 0xE3, 0xC4), False, R)], "sa": 2, "line": 1.15},
      {"runs": [("Arithmetic where known.", 12.5, True, RGBColor(0x9B, 0xE3, 0xC4), False, R)], "sa": 3, "line": 1.15},
      {"runs": [("The LLM never decides the money.", 9.5, False, SOFTB, False, S)], "line": 1.15}])

text(s4, 0.54, 2.62, 5.6, 0.26, [{"runs": [("TWO HALVES, DELIBERATELY SEPARATED  ·  PLUS A SOLVER", 9.5, True, BLUE, False, R)]}])
for i, (t1, t2, col) in enumerate([
        ("A · FORECAST", "Predicts the itemised bill from published tariff data. Uncertain, so it returns a range.", BLUE),
        ("B · ADJUDICATION", "Computes what the policy refuses, line by line. Exact, with every ₹ cited to its clause.", GREEN),
        ("C · COUNTERFACTUAL", "Re-runs the claim for every legal alternative and ranks them by what you keep.", BLUE)]):
    yy = 2.86 + i * 0.88
    rect(s4, 0.54, yy, 0.07, 0.80, col)
    text(s4, 0.86, yy, 5.3, 0.28, [{"runs": [(t1, 14, True, col, False, R)]}])
    text(s4, 0.86, yy + 0.30, 5.3, 0.50, [{"runs": [(t2, 11.5, False, INK, False, S)], "line": 1.3}])

RX, RW = 6.86, 5.93
text(s4, RX, 2.62, RW, 0.26, [{"runs": [("AGAINST THE CONVENTIONAL ALTERNATIVES", 9.5, True, BLUE, False, R)]}])
rows = [("", "Hospital / TPA desk", "Policy bots · Ditto/PB", "ClaimCast"),
        ("Speaks", "after admission", "any time, no context", "at each decision"),
        ("Returns", "a verbal estimate", "a text summary", "itemised ₹ + clause"),
        ("Acts for", "the hospital", "no one", "you")]
cw = [0.94, 1.64, 1.70, 1.65]
ry = 2.94
for ri, row in enumerate(rows):
    cx = RX
    for ci, cell in enumerate(row):
        head, last = ri == 0, ci == 3
        col = BLUE if head else (INK if last else BODY)
        text(s4, cx, ry, cw[ci] - 0.06, 0.28,
             [{"runs": [(cell, 10, True, col, False, S) if (head or last) else (cell, 10.5, False, col, False, S)]}])
        cx += cw[ci]
    ry += 0.28 if ri == 0 else 0.36
    hline(s4, RX, ry - 0.09, RW, HAIR)
pc = ry + 0.06
text(s4, RX, pc, 2.7, 0.24, [{"runs": [("PROS", 9.5, True, GREEN, False, R)]}])
text(s4, RX, pc + 0.24, 2.76, 0.78,
     [{"runs": [("•  Deterministic and auditable", 10.5, False, INK, False, S)], "sa": 3},
      {"runs": [("•  Runs entirely on public data", 10.5, False, INK, False, S)], "sa": 3},
      {"runs": [("•  Rule-packs: private, PM-JAY, ESI", 10.5, False, INK, False, S)]}])
text(s4, RX + 3.10, pc, 2.7, 0.24, [{"runs": [("CONS", 9.5, True, RED, False, R)]}])
text(s4, RX + 3.10, pc + 0.24, 2.80, 0.78,
     [{"runs": [("•  Forecast bounded by published tariffs", 10.5, False, INK, False, S)], "sa": 3},
      {"runs": [("•  Needs the policy document up front", 10.5, False, INK, False, S)], "sa": 3},
      {"runs": [("•  An estimate, never a guarantee", 10.5, False, INK, False, S)]}])

SCY, SCX, SCW = 5.50, 0.54, W - 1.08
rect(s4, SCX, SCY, SCW, 1.24, WHITE, RGBColor(0xC9, 0xD8, 0xE6), 0.75)
rect(s4, SCX, SCY, SCW, 0.30, BLUE)
text(s4, SCX + 0.22, SCY + 0.07, 8.0, 0.22,
     [{"runs": [("ClaimCast", 9.5, True, WHITE, False, R),
                ("      DECISION 2 OF 5   ·   ROOM CLASS   ·   STAGE: ADMISSION",
                 9, False, PALEB, False, S)]}])
text(s4, SCX + SCW - 2.68, SCY + 0.08, 2.46, 0.22,
     [{"runs": [("PROTOTYPE SCREEN", 8, True, RGBColor(0x9E, 0xC1, 0xE4), False, S)],
       "align": PP_ALIGN.RIGHT}])
for i, (mark, lab, rate, pay, col) in enumerate([
        ("○", "Private room", "₹10,000 / day", "₹1,26,900", RED),
        ("●", "Semi-private", "₹5,000 / day", "₹48,500", BLUE)]):
    yy = SCY + 0.40 + i * 0.25
    text(s4, SCX + 0.22, yy, 0.22, 0.22, [{"runs": [(mark, 10, False, col, False, S)]}])
    text(s4, SCX + 0.52, yy, 2.4, 0.22, [{"runs": [(lab, 11, False, INK, False, S)]}])
    text(s4, SCX + 3.00, yy, 1.6, 0.22, [{"runs": [(rate, 10, False, BODY, False, S)]}])
    text(s4, SCX + 4.80, yy - 0.02, 2.8, 0.24,
         [{"runs": [("your share  ", 9.5, False, BODY, False, S), (pay, 12, True, col, False, R)]}])
    if i:
        text(s4, SCX + 7.80, yy, 3.6, 0.22,
             [{"runs": [("←  recommended by ClaimCast", 9.5, False, GREEN, False, S)]}])
hline(s4, SCX + 0.22, SCY + 0.94, SCW - 0.44, HAIR)
text(s4, SCX + 0.22, SCY + 1.00, 11.6, 0.24,
     [{"runs": [("✓  You keep ₹78,400", 13, True, GREEN, False, R),
                ("      Room Rent Sub-limit, Clause 3.2  ·  proportionate deduction on the "
                 "associated charges", 9.5, False, BODY, False, S)]}])

# ══════════════════════════════════════════════════════════════ 5 · REFERENCES
retitle(s5, "Any Additional Information / References", font=R, color=TITLE)
content_frame(s5, keep=False)
text(s5, 0.54, 1.20, 5.6, 0.26, [{"runs": [("ARCHITECTURE", 9.5, True, BLUE, False, R)]}])
AXX, AWW = 0.54, 5.72
half = (AWW - 0.18) / 2


def pbox(x, y, w, h, t1, t2, fill, tcol, scol=BODY):
    rect(s5, x, y, w, h, fill)
    blocks = [{"runs": [(t1, 11.5, True, tcol, False, R)], "align": PP_ALIGN.CENTER, "sa": 1}]
    if t2:
        blocks.append({"runs": [(t2, 8.5, False, scol, False, S)], "align": PP_ALIGN.CENTER})
    text(s5, x + 0.06, y + (0.11 if t2 else 0.16), w - 0.12, h - 0.16, blocks)


TB = RGBColor(0xEC, 0xF2, 0xF8)
pbox(AXX, 1.50, half, 0.52, "Policy PDF", "insurer wording", TB, INK)
pbox(AXX + half + 0.18, 1.50, half, 0.52, "Procedure · Hospital · City", "the care plan", TB, INK)
arrow(s5, AXX + half / 2 - 0.07, 2.06, 0.14, 0.14, MSO_SHAPE.DOWN_ARROW)
arrow(s5, AXX + half + 0.18 + half / 2 - 0.07, 2.06, 0.14, 0.14, MSO_SHAPE.DOWN_ARROW)
pbox(AXX, 2.26, half, 0.56, "LLM extract + cite", "user confirms every field", TB, BLUE)
pbox(AXX + half + 0.18, 2.26, half, 0.56, "ML cost model", "NHA / CGHS tariffs", TB, BLUE)
arrow(s5, AXX + half / 2 - 0.07, 2.86, 0.14, 0.14, MSO_SHAPE.DOWN_ARROW)
arrow(s5, AXX + half + 0.18 + half / 2 - 0.07, 2.86, 0.14, 0.14, MSO_SHAPE.DOWN_ARROW)
pbox(AXX, 3.06, AWW, 0.58, "ADJUDICATION ENGINE  ·  deterministic",
     "every ₹ refused, cited to its clause", GREEN, WHITE, RGBColor(0xC5, 0xE3, 0xD4))
arrow(s5, AXX + AWW / 2 - 0.07, 3.68, 0.14, 0.14, MSO_SHAPE.DOWN_ARROW)
pbox(AXX, 3.88, AWW, 0.56, "COUNTERFACTUAL SOLVER", "cheapest fully-covered path", TB, BLUE)
arrow(s5, AXX + AWW / 2 - 0.07, 4.48, 0.14, 0.14, MSO_SHAPE.DOWN_ARROW)
pbox(AXX, 4.68, AWW, 0.54, "5 DECISION SCREENS  ·  filtered by materiality", None, BLUE, WHITE)

text(s5, RX, 1.20, RW, 0.26,
     [{"runs": [("WHAT WE BUILD, AND WHAT WE SIMULATE", 9.5, True, BLUE, False, R)]}])
rm = [("Policy → constraint parser", "Build", True), ("Adjudication + deduction engine", "Build", True),
      ("Cost model on NHA / CGHS rates", "Build", True), ("Counterfactual solver + screens", "Build", True),
      ("Live hospital tariff feed", "Simulate", False),
      ("Insurer cashless API / HIS bill feed", "Simulate", False)]
ry = 1.50
for what, state, real in rm:
    rect(s5, RX, ry + 0.05, 0.06, 0.17, GREEN if real else PALEB)
    text(s5, RX + 0.20, ry, 4.3, 0.26, [{"runs": [(what, 10.5, False, INK, False, S)]}])
    text(s5, RX + 4.52, ry, 1.4, 0.26, [{"runs": [(state, 10.5, True, GREEN if real else BLUE, False, R)]}])
    ry += 0.32
    hline(s5, RX, ry - 0.07, RW, HAIR)
text(s5, RX, ry + 0.06, RW, 0.24,
     [{"runs": [("The two simulated feeds need an insurer and a hospital to open them, "
                 "not more engineering.", 9.5, False, BODY, False, S)]}])
ds = ry + 0.38
text(s5, RX, ds, RW, 0.26, [{"runs": [("DATA WE SOURCE OURSELVES  ·  ALL PUBLIC", 9.5, True, BLUE, False, R)]}])
text(s5, RX, ds + 0.28, RW, 0.58,
     [{"runs": [("NHA / PM-JAY Health Benefit Packages (≈1,900 procedure rates)  ·  CGHS city-wise "
                 "rate lists  ·  IRDAI List I non-payables  ·  IRDAI standard wordings (Arogya "
                 "Sanjeevani)  ·  NHA empanelled-hospital registry", 10, False, BODY, False, S)], "line": 1.3}])
rf = ds + 0.98
text(s5, RX, rf, RW, 0.26, [{"runs": [("REFERENCES", 9.5, True, BLUE, False, R)]}])
for i, r_ in enumerate([
        "1.  IRDAI, Circular IRDAI/HLT/REG/CIR/151/06/2020, room-rent proportionate deduction.",
        "2.  Financial toxicity of cancer treatment in India, Front. Public Health, 2023 (n = 12,148).",
        "3.  Catastrophic expenditure & treatment attrition, Lancet Reg. Health SE Asia, 2022;6:100058.",
        "4.  IRDAI Master Circular on Health Insurance Business, 2024."]):
    text(s5, RX, rf + 0.28 + i * 0.22, RW, 0.22, [{"runs": [(r_, 9.5, False, BODY, False, S)]}])

hline(s5, 0.54, 6.02, W - 1.08, RULE)
text(s5, 0.54, 6.18, 12.25, 0.32,
     [{"runs": [("SAFE BY CONSTRUCTION    ", 9.5, True, GREEN, False, R),
                ("every field is user-confirmed  ·  outputs are labelled estimates, never "
                 "guarantees  ·  no medical advice, money only  ·  no identifiable patient data",
                 10.5, False, BODY, False, S)]}])

prs.save(OUT)
print("saved:", OUT)
