# -*- coding: utf-8 -*-
"""Shared kit for building PCC-2026 decks on the official GE template."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import os as _os
import glob as _glob
_HERE = _os.path.dirname(_os.path.abspath(__file__))
ROOT = _os.path.dirname(_HERE)
# the official template, wherever it ends up living under the project
_hits = _glob.glob(_os.path.join(ROOT, "**", "PCC 2026 TEMPLATE.pptx"), recursive=True)
SRC = _hits[0] if _hits else _os.path.join(_HERE, "PCC 2026 TEMPLATE.pptx")

BLUE  = RGBColor(0x00, 0x5E, 0xB8)
PALE  = RGBColor(0xA9, 0xCB, 0xEA)
INK   = RGBColor(0x1B, 0x1B, 0x1B)
BODY  = RGBColor(0x55, 0x58, 0x5A)
MUTED = RGBColor(0x9A, 0x9C, 0x9E)
GREY  = RGBColor(0x63, 0x66, 0x6A)
RED   = RGBColor(0xC8, 0x10, 0x2E)
GREEN = RGBColor(0x00, 0x77, 0x4F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE  = RGBColor(0xD8, 0xD4, 0xCE)
HAIR  = RGBColor(0xDC, 0xE0, 0xE4)
BLACK = RGBColor(0x1B, 0x1B, 0x1B)

TITLE_X, TITLE_Y = 0.54, 0.28
FX, FY, FW, FH = 0.54, 1.22, 12.25, 5.62      # template content frame
IX, IY, IW = 0.85, 1.48, 11.63                 # padded inner area
IB = 6.58                                      # inner bottom

_font = {"display": "Georgia", "body": "Segoe UI"}


def set_fonts(display, body):
    _font["display"], _font["body"] = display, body


def D():
    return _font["display"]


def B():
    return _font["body"]


def para(tf, first, runs, align=PP_ALIGN.LEFT, sb=0, sa=0, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(sb)
    p.space_after = Pt(sa)
    if line:
        p.line_spacing = line
    for r_ in runs:
        t, sz, b, c = r_[0], r_[1], r_[2], r_[3]
        italic = r_[4] if len(r_) > 4 else False
        fname = r_[5] if len(r_) > 5 else _font["body"]
        run = p.add_run()
        run.text = t
        f = run.font
        f.name, f.size, f.bold, f.color.rgb, f.italic = fname, Pt(sz), b, c, italic
    return p


def text(slide, x, y, w, h, blocks, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, blk in enumerate(blocks):
        para(tf, i == 0, blk["runs"], blk.get("align", PP_ALIGN.LEFT),
             blk.get("sb", 0), blk.get("sa", 0), blk.get("line"))
    return tb


def rect(slide, x, y, w, h, fill, line=None, lw=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def hline(slide, x, y, w, color=RULE, weight=0.012):
    return rect(slide, x, y, w, weight, color)


def vline(slide, x, y, h, color=RULE, weight=0.012):
    return rect(slide, x, y, weight, h, color)


def arrow(slide, x, y, w, h, direction=MSO_SHAPE.RIGHT_ARROW, color=None):
    s = slide.shapes.add_shape(direction, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color or RGBColor(0xC3, 0xD3, 0xE3)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def shape_named(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def drop(slide, *names):
    for n in names:
        for sh in list(slide.shapes):
            if sh.name == n:
                sh._element.getparent().remove(sh._element)


def send_to_back(shape):
    el = shape._element
    parent = el.getparent()
    parent.remove(el)
    parent.insert(2, el)


def retitle(slide, txt, size=30, color=GREY, font=None, bold=True):
    """Restyle the template's slide-title text box."""
    sh = shape_named(slide, "TextShape 2")
    sh.left, sh.top = Inches(TITLE_X), Inches(TITLE_Y)
    sh.width, sh.height = Inches(11.9), Inches(0.62)
    tf = sh.text_frame
    tf.clear()
    para(tf, True, [(txt, size, bold, color, False, font or _font["display"])])
    return sh


def content_frame(slide, keep=True, fill=WHITE, line=BLUE):
    """Resize/clear the template's bordered content box; keep=False removes it."""
    sh = shape_named(slide, "TextShape 1")
    if sh is None:
        return None
    if not keep:
        sh._element.getparent().remove(sh._element)
        return None
    sh.left, sh.top = Inches(FX), Inches(FY)
    sh.width, sh.height = Inches(FW), Inches(FH)
    tf = sh.text_frame
    tf.clear()
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def open_deck():
    prs = Presentation(SRC)
    s = list(prs.slides)
    drop(s[2], "CustomShape 5", "CustomShape 6")     # template's leftover prompt shapes
    return prs, s


def title_slide(s1, product, tagline_a, tagline_b, ps_line, team, college,
                accent=BLUE, display=None, body=None):
    """Common title-slide skeleton; directions restyle around it."""
    disp = display or _font["display"]
    bod = body or _font["body"]
    t = s1.shapes.title
    t.left, t.top, t.width, t.height = Inches(1.0), Inches(0.62), Inches(11.33), Inches(0.62)
    tf = t.text_frame
    tf.clear()
    para(tf, True, [("Precision Care Challenge 2026", 20, False, GREY, False, bod)], PP_ALIGN.CENTER)

    sub = [p for p in s1.placeholders if p.placeholder_format.idx == 1][0]
    sub.left, sub.top, sub.width, sub.height = Inches(1.0), Inches(1.24), Inches(11.33), Inches(0.36)
    tf = sub.text_frame
    tf.clear()
    para(tf, True, [("Phase 1  \u2013  Idea Submission", 12, False, MUTED, False, bod)], PP_ALIGN.CENTER)
    return disp, bod
